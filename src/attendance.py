# -*- coding: utf-8 -*-
"""회사 월 근무현황 엑셀에서 특정 인원의 출퇴근 기록을 뽑는다.

시트: 월별 1장. 3행에 성명이 4칸 간격으로, 4행에 주/야·시작·종료 머리글,
5행부터 날짜별 기록. 빈 줄은 출근하지 않은 날(주말·공휴일·휴가 구분 없음).
"""
import sys
from datetime import date, datetime, time

WORKBOOK = (r"C:/Users/willi/OneDrive/Desktop/AZtech_macro_maker"
            r"/작업자 목록/02. 26년도 ○○실 월 근무현황.xlsx")


def _as_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def _as_time(v):
    if isinstance(v, datetime):
        return v.time()
    if isinstance(v, time):
        return v
    return None


def find_column(ws, name, header_row=3):
    for c in range(1, ws.max_column + 1):
        if str(ws.cell(header_row, c).value or "").strip() == name:
            return c
    return None


def read_person(path, name):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = []
    for ws in wb.worksheets:
        col = find_column(ws, name)
        if col is None:
            continue
        for r in range(5, ws.max_row + 1):
            d = _as_date(ws.cell(r, 1).value)
            if d is None:
                continue
            shift = ws.cell(r, col).value
            start = _as_time(ws.cell(r, col + 1).value)
            end = _as_time(ws.cell(r, col + 2).value)
            note = ws.cell(r, col + 3).value
            rows.append({
                "date": d,
                "shift": str(shift).strip() if shift else None,
                "start": start,
                "end": end,
                "note": str(note).strip() if note else None,
                "worked": bool(start and end),
            })
    rows.sort(key=lambda x: x["date"])
    return rows


LUNCH_START, LUNCH_END = 12.0, 13.0


def worked_hours(row):
    """출퇴근 시각에서 실근무 시간. 점심시간(12~13시)과 겹치는 만큼만 뺀다.
    오후 반차처럼 13시 출근인 날은 점심을 빼면 안 된다."""
    if not row["worked"]:
        return 0.0
    s, e = row["start"], row["end"]
    sh = s.hour + s.minute / 60
    eh = e.hour + e.minute / 60
    if eh < sh:              # 야간 교대
        eh += 24
    overlap = max(0.0, min(eh, LUNCH_END) - max(sh, LUNCH_START))
    return round(eh - sh - overlap, 2)
