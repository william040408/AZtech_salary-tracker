# -*- coding: utf-8 -*-
"""회사 근무현황(엑셀)과 급여명세서를 대조한다.

근무현황의 빈칸은 두 종류다. 그 날 전원이 비어 있으면 회사 전체 휴무이고,
본인만 비어 있으면 개인 휴가다. 명세서의 기본급 차감과 대조할 대상은
후자 쪽이므로 둘을 구분해서 센다.
"""
import sys
from datetime import date, datetime
from collections import defaultdict

sys.path.insert(0, "src")
from attendance import find_column, _as_date, _as_time, worked_hours, WORKBOOK
from calendar_kr import holidays

STD_HOURS = 8.0


def read_all(path):
    """{날짜: {이름: row}} 전원 기록."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    days = defaultdict(dict)
    for ws in wb.worksheets:
        names = []
        for c in range(1, ws.max_column + 1):
            n = str(ws.cell(3, c).value or "").strip()
            if n and n not in ("일자 성명", "확인"):
                names.append((c, n))
        for r in range(5, ws.max_row + 1):
            d = _as_date(ws.cell(r, 1).value)
            if d is None:
                continue
            for c, n in names:
                st, en = _as_time(ws.cell(r, c + 1).value), _as_time(ws.cell(r, c + 2).value)
                days[d][n] = {"date": d, "start": st, "end": en,
                              "worked": bool(st and en),
                              "note": str(ws.cell(r, c + 3).value or "").strip() or None}
    return days


def classify(days, me, first_day, last_day):
    """월별로 회사휴무 / 개인휴가 / 반차 / 지각을 집계."""
    H = holidays(2026)
    out = defaultdict(lambda: {"company_off": [], "personal_off": [], "half": [],
                               "late_hours": 0.0, "worked": 0, "hours": 0.0})
    for d in sorted(days):
        if d < first_day or d > last_day:
            continue
        if d.weekday() >= 5 or d in H:
            continue
        roster = days[d]
        if me not in roster:
            continue
        m = d.strftime("%Y-%m")
        anyone = any(v["worked"] for v in roster.values())
        row = roster[me]
        if not row["worked"]:
            (out[m]["company_off"] if not anyone else out[m]["personal_off"]).append(d)
            continue
        h = worked_hours(row)
        out[m]["worked"] += 1
        out[m]["hours"] += h
        if h <= 4.6:
            out[m]["half"].append((d, h))
        elif h < STD_HOURS:
            out[m]["late_hours"] += STD_HOURS - h
    return out
